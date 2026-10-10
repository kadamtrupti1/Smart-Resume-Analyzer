# parser/resume_parser.py
# Handles PDF and DOCX resume text extraction + section detection.
# This module focuses on EXTRACTING text and FINDING resume sections.
# It does NOT score or analyze — that is done by the analyzer modules.

import re
import os

# ---------------------------------------------------------------------------
# PDF EXTRACTION
# We use pdfplumber as the primary library (it handles tables and formatting).
# PyPDF2 is used as a fallback if pdfplumber fails.
# ---------------------------------------------------------------------------

def extract_text_from_pdf(file_path):
    """
    Extract all text from a PDF file using pdfplumber.
    Falls back to PyPDF2 if pdfplumber fails.

    Returns: A single string containing all the resume text.
    """
    text = ""

    # Try pdfplumber first
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        if text.strip():
            return text
    except Exception as e:
        print(f"[pdfplumber] Failed: {e}. Trying PyPDF2...")

    # Fallback: try PyPDF2
    try:
        import PyPDF2
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        print(f"[PyPDF2] Failed: {e}")

    return text


# ---------------------------------------------------------------------------
# DOCX EXTRACTION
# We use python-docx to extract paragraphs from Word documents.
# ---------------------------------------------------------------------------

def extract_text_from_docx(file_path):
    """
    Extract all text from a DOCX file using python-docx.

    Returns: A single string containing all the resume text.
    """
    text = ""
    try:
        from docx import Document
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"

        # Also extract text from tables (some resumes use tables for layout)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"
    except Exception as e:
        print(f"[python-docx] Failed: {e}")

    return text


# ---------------------------------------------------------------------------
# SECTION DETECTION
# We look for common section headings in the resume text.
# This uses simple string matching — no AI needed.
# ---------------------------------------------------------------------------

# Map of section names to possible heading keywords
SECTION_KEYWORDS = {
    'objective':       ['objective', 'career objective', 'professional objective'],
    'summary':         ['summary', 'professional summary', 'profile', 'about me', 'about'],
    'education':       ['education', 'academic background', 'academic qualification', 'qualifications'],
    'experience':      ['experience', 'work experience', 'professional experience', 'employment', 'internship', 'internships'],
    'skills':          ['skills', 'technical skills', 'core skills', 'key skills', 'technologies', 'technical expertise', 'competencies'],
    'projects':        ['projects', 'academic projects', 'personal projects', 'key projects', 'project work'],
    'certifications':  ['certifications', 'certification', 'certificates', 'courses', 'online courses', 'achievements'],
    'contact':         ['contact', 'contact information', 'contact details', 'personal details', 'personal information'],
}


def detect_sections(text):
    """
    Scan the resume text line by line.
    Identify which sections are present based on heading keywords.

    Returns a dict:
    {
        'section_name': 'text content of that section',
        ...
    }
    """
    lines = text.split('\n')
    sections = {}
    current_section = 'header'
    sections[current_section] = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Check if this line is a section heading
        found_section = _identify_heading(stripped)

        if found_section:
            current_section = found_section
            if current_section not in sections:
                sections[current_section] = []
        else:
            sections[current_section].append(stripped)

    # Convert lists to strings
    result = {}
    for section, lines_list in sections.items():
        result[section] = '\n'.join(lines_list).strip()

    return result


def _identify_heading(line):
    """
    Check if a given line matches any known section heading.
    Returns the section name if found, else None.
    """
    # Normalize: lowercase, strip punctuation
    normalized = re.sub(r'[^a-z\s]', '', line.lower()).strip()

    for section_name, keywords in SECTION_KEYWORDS.items():
        for keyword in keywords:
            # Exact match or the line starts with the keyword
            if normalized == keyword or normalized.startswith(keyword):
                return section_name

    return None


# ---------------------------------------------------------------------------
# MAIN PARSER FUNCTION
# This is the function called by app.py
# ---------------------------------------------------------------------------

def parse_resume(file_path):
    """
    Main function: Parse a resume file (PDF or DOCX).

    Steps:
    1. Detect file type
    2. Extract text
    3. Detect sections
    4. Extract contact info (email, phone, name)

    Returns a dict with:
    - 'raw_text': full extracted text
    - 'sections': detected sections dict
    - 'name': candidate name
    - 'email': candidate email
    - 'phone': candidate phone
    - 'sections_found': list of section names found
    """
    from utils.helpers import clean_text, extract_email, extract_phone, extract_name_from_text

    # Step 1: Detect file type and extract text
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        raw_text = extract_text_from_pdf(file_path)
    elif ext == '.docx':
        raw_text = extract_text_from_docx(file_path)
    else:
        return {'error': 'Unsupported file type'}

    if not raw_text.strip():
        return {'error': 'Could not extract text from the file. Please check the file is not image-only.'}

    # Step 2: Clean the text
    clean = clean_text(raw_text)

    # Step 3: Detect sections
    sections = detect_sections(clean)

    # Step 4: Extract basic contact info
    email = extract_email(clean)
    phone = extract_phone(clean)
    name  = extract_name_from_text(clean)

    # Determine which key sections were found
    key_sections = ['education', 'experience', 'skills', 'projects', 'certifications', 'summary', 'objective']
    sections_found = [s for s in key_sections if s in sections and sections[s]]

    return {
        'raw_text':      clean,
        'sections':      sections,
        'name':          name,
        'email':         email,
        'phone':         phone,
        'sections_found': sections_found,
        'word_count':    len(clean.split()),
    }
